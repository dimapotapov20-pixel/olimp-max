#!/usr/bin/env python3
"""Crawl official admissions targets and send keyword-based candidates to PostgreSQL.

The crawler is deliberately source-agnostic.  It reads active, due targets
from ``admission_source_targets``, saves a content-addressed source snapshot,
uses Requests + BeautifulSoup for HTML and small native readers for PDF/XLSX,
then creates review candidates from tables and paragraphs around admissions
keywords.

It never writes to ``benefit_rules``. A reviewer maps a candidate to a real
programme and olympiad profile before publishing it.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import UTC, datetime
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

try:
    from parser.strict_admission_adapters import StrictRule, supported_codes as strict_adapter_codes, strict_rules
except ModuleNotFoundError:  # Allows `python parser/crawl_admission_rules.py`.
    from strict_admission_adapters import StrictRule, supported_codes as strict_adapter_codes, strict_rules

try:
    import requests
except ImportError:  # pragma: no cover - reported by main
    requests = None

try:
    import psycopg
    from psycopg.rows import dict_row
    from psycopg.types.json import Jsonb
except ImportError:  # pragma: no cover - reported by main
    psycopg = None
    dict_row = None
    Jsonb = None


ROOT = Path(__file__).resolve().parents[1]
SNAPSHOT_DIR = ROOT / "parser" / "snapshots"
ADAPTER_CODE = "requests-bs4-keyword-v1"
ADAPTER_VERSION = "2026.8"
# Lets an existing local database transition without a manual data update.
SUPPORTED_TARGET_CODES = (ADAPTER_CODE, "docling-keyword-v1", *strict_adapter_codes())
USER_AGENT = "OlimpAdmissionCrawler/0.1 (+educational-demo; public-source-reader)"
MAX_DOCUMENT_BYTES = 30 * 1024 * 1024

BENEFIT_TERMS = {
    "bvi": ("бви", "без вступительных испытаний"),
    "hundred_points": ("100 баллов", "100 балл"),
}
CONTEXT_TERMS = (
    "олимпиад",
    "особое право",
    "победител",
    "призер",
    "вступительных испытаний",
    "егэ",
)
TABLE_HEADER_TERMS = (
    "наименование олимпиады",
    "название олимпиады",
    "профиль олимпиады",
    "перечень олимпиад",
    "уровень олимпиады",
    "предоставляемая льгота",
    "право на прием",
    "право на 100 баллов",
    "поступление бви",
)
SUBJECT_CODES = (
    ("информат", "informatics"), ("программир", "informatics"),
    ("математ", "math"), ("физик", "physics"), ("хими", "chemistry"),
    ("биолог", "biology"), ("астроном", "astronomy"), ("эконом", "economics"),
    ("истори", "history"), ("географ", "geography"), ("обществ", "social-studies"),
    ("прав", "law"), ("литератур", "literature"), ("русск", "russian"),
    ("англий", "english"), ("немец", "german"), ("француз", "french"),
    ("китай", "chinese"),
)


@dataclass(frozen=True)
class Target:
    id: int
    source_id: int
    campaign_id: int
    university_location_id: int | None
    adapter_code: str
    url: str
    source_base_url: str
    adapter_config: dict[str, Any]


@dataclass(frozen=True)
class KeywordCandidate:
    candidate_key: str
    source_locator: str
    source_excerpt: str
    raw_rule_text: str
    suggested_benefit_kind: str | None
    suggested_diploma_status: str | None
    suggested_confirmation_subject_code: str | None
    suggested_confirmation_min_score: int | None
    confidence: int
    raw_payload: dict[str, Any]
    raw_olympiad_name: str | None = None
    raw_profile_name: str | None = None
    raw_programme_name: str | None = None
    suggested_point_value: int | None = None
    suggested_olympiad_level: int | None = None


def compact(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00ad", "").replace("\xa0", " ")).strip()


def normalized(value: str) -> str:
    return compact(value).casefold().replace("ё", "е")


def stable_key(*parts: str) -> str:
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:32]


def adapter_version(target: Target) -> str:
    """A target changing parser mode must reprocess an unchanged document."""
    return f"{target.adapter_code}:{ADAPTER_VERSION}"


def guess_suffix(content_type: str, url: str) -> str:
    content_type = content_type.split(";", 1)[0].strip().lower()
    by_type = {
        "text/html": ".html", "application/xhtml+xml": ".html", "application/pdf": ".pdf",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
        "application/vnd.ms-excel": ".xls",
    }
    if content_type in by_type:
        return by_type[content_type]
    suffix = Path(urlparse(url).path).suffix.lower()
    return suffix if suffix in {".html", ".htm", ".pdf", ".xlsx", ".xls"} else ".bin"


def parsed_http_date(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return parsedate_to_datetime(value)
    except (TypeError, ValueError):
        return None


def allowed_host(url: str, base_url: str, configured_hosts: list[str]) -> bool:
    host = (urlparse(url).hostname or "").lower()
    base_host = (urlparse(base_url).hostname or "").lower()
    allowed = {base_host, *(item.lower() for item in configured_hosts)}
    return bool(host and any(host == item or host.endswith(f".{item}") for item in allowed))


def fetch(target: Target) -> tuple[bytes, str, str, str | None, str | None]:
    if requests is None:
        raise RuntimeError("requests is not installed. Run: python3 -m pip install -r requirements-crawler.txt")
    allowed_hosts = list(target.adapter_config.get("allow_hosts", []))
    user_agent = target.adapter_config.get("user_agent", USER_AGENT)
    if not isinstance(user_agent, str) or not user_agent.strip() or len(user_agent) > 240:
        raise ValueError("Configured User-Agent must be a non-empty string of at most 240 characters")
    if urlparse(target.url).scheme != "https" or not allowed_host(target.url, target.source_base_url, allowed_hosts):
        raise ValueError("Target URL must be HTTPS and use an approved official host")
    response = requests.get(
        target.url,
        allow_redirects=True,
        timeout=(5, 30),
        headers={
            "User-Agent": user_agent,
            "Accept": "text/html,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;q=0.9,*/*;q=0.1",
        },
    )
    response.raise_for_status()
    final_url = str(response.url)
    if not allowed_host(final_url, target.source_base_url, allowed_hosts):
        raise ValueError("Official source redirected to an unapproved host")
    if len(response.content) > MAX_DOCUMENT_BYTES:
        raise ValueError(f"Document exceeds {MAX_DOCUMENT_BYTES // 1024 // 1024} MB limit")
    return (
        response.content,
        final_url,
        response.headers.get("content-type", "application/octet-stream"),
        response.headers.get("etag"),
        response.headers.get("last-modified"),
    )


def html_to_markdown(snapshot: Path) -> str:
    """Keep headings, paragraphs and tables from an official HTML page."""
    try:
        from bs4 import BeautifulSoup
    except ImportError as error:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "BeautifulSoup is not installed. Run: python3 -m pip install -r requirements-crawler.txt"
        ) from error
    soup = BeautifulSoup(snapshot.read_bytes(), "html.parser")
    for node in soup(["script", "style", "noscript", "svg", "template"]):
        node.decompose()

    lines: list[str] = []
    root = soup.body or soup
    for node in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr"]):
        if node.name != "tr" and node.find_parent("table"):
            continue
        if node.name == "tr":
            cells = [
                compact(cell.get_text(" ", strip=True))
                for cell in node.find_all(["th", "td"], recursive=False)
            ]
            if not any(cells):
                continue
            lines.append("| " + " | ".join(cells) + " |")
            if node.find("th", recursive=False):
                lines.append("| " + " | ".join("---" for _ in cells) + " |")
            continue
        text = compact(node.get_text(" ", strip=True))
        if not text:
            continue
        if node.name.startswith("h"):
            lines.append("#" * int(node.name[1]) + " " + text)
        elif node.name == "li":
            lines.append("- " + text)
        else:
            lines.append(text)
    return "\n".join(lines)


def pdf_to_markdown(snapshot: Path) -> str:
    """Read digital PDF text and detectable tables without OCR or ML models."""
    try:
        import pdfplumber
    except ImportError as error:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "pdfplumber is not installed. Run: python3 -m pip install -r requirements-crawler.txt"
        ) from error
    lines: list[str] = []
    with pdfplumber.open(snapshot) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            lines.append(f"## Страница {page_number}")
            text = page.extract_text() or ""
            lines.extend(compact(line) for line in text.splitlines() if compact(line))
            for table in page.extract_tables():
                for row_index, row in enumerate(table):
                    cells = [compact(cell or "") for cell in row]
                    if any(cells):
                        lines.append("| " + " | ".join(cells) + " |")
                        if row_index == 0:
                            lines.append("| " + " | ".join("---" for _ in cells) + " |")
    return "\n".join(lines)


def xlsx_to_markdown(snapshot: Path) -> str:
    """Read spreadsheet cells while retaining rows for the keyword extractor."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "openpyxl is not installed. Run: python3 -m pip install -r requirements-crawler.txt"
        ) from error
    lines: list[str] = []
    workbook = load_workbook(snapshot, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            lines.append(f"## Лист: {worksheet.title}")
            has_header = False
            for row in worksheet.iter_rows(values_only=True):
                cells = [compact("" if value is None else str(value)) for value in row]
                if not any(cells):
                    continue
                lines.append("| " + " | ".join(cells) + " |")
                if not has_header:
                    lines.append("| " + " | ".join("---" for _ in cells) + " |")
                    has_header = True
    finally:
        workbook.close()
    return "\n".join(lines)


def document_to_markdown(snapshot: Path) -> str:
    suffix = snapshot.suffix.lower()
    if suffix in {".html", ".htm"}:
        return html_to_markdown(snapshot)
    if suffix == ".pdf":
        return pdf_to_markdown(snapshot)
    if suffix in {".xlsx", ".xls"}:
        return xlsx_to_markdown(snapshot)
    raise ValueError(f"Unsupported snapshot format: {snapshot.suffix}")


def inferred_benefit(value: str) -> str | None:
    value = normalized(value)
    for kind, terms in BENEFIT_TERMS.items():
        if any(term in value for term in terms):
            return kind
    return None


def inferred_diploma(value: str) -> str | None:
    value = normalized(value)
    winner = "победител" in value
    prize_winner = "призер" in value
    if winner and not prize_winner:
        return "winner"
    if prize_winner and not winner:
        return "prize_winner"
    return None  # "winners and prize winners" must be resolved by a human.


def inferred_subject(value: str) -> str | None:
    value = normalized(value)
    for fragment, code in SUBJECT_CODES:
        if fragment in value:
            return code
    return None


def inferred_min_score(value: str) -> int | None:
    value = normalized(value)
    match = re.search(r"(?:егэ|ви)[^0-9]{0,56}\b(\d{2,3})\s*бал", value)
    if not match:
        return None
    score = int(match.group(1))
    return score if 0 <= score <= 100 else None


def markdown_cells(row: str) -> list[str]:
    return [compact(cell) for cell in row.strip().strip("|").split("|")]


def olympiad_title_column_index(header_cells: list[str]) -> int | None:
    return next(
        (
            index
            for index, cell in enumerate(header_cells)
            if cell in {"олимпиада", "наименование олимпиады", "название олимпиады"}
            or "наименование олимпиады" in cell
            or "название олимпиады" in cell
        ),
        None,
    )


def explicit_row_olympiad_title(headers: list[str], row: str) -> str | None:
    """Read only a dedicated title column; list columns are handled separately."""
    header_line = next((item for item in headers if "|" in item), None)
    if header_line is None:
        return None
    title_index = olympiad_title_column_index([normalized(cell) for cell in markdown_cells(header_line)])
    row_cells = markdown_cells(row)
    if title_index is None or len(row_cells) <= title_index:
        return None
    return row_cells[title_index] or None


def explicit_olympiad_names(value: str) -> list[str]:
    """Split a named olympiad list, but never guess from ordinary prose."""
    markers = (
        "олимпиада",
        "всероссийская",
        "межрегиональная",
        "международная",
        "московская",
        "инженерная",
        "интернет",
        "открытая",
        "вузовско",
        "многопрофильная",
        "городская",
        "всесибирская",
        "междисциплинарная",
    )
    pattern = r"(?:,|;)\s+(?=(?:" + "|".join(markers) + r")\b)"
    return [compact(item) for item in re.split(pattern, value, flags=re.IGNORECASE) if compact(item)]


def table_identities(payload: dict[str, Any]) -> list[dict[str, str]]:
    """Extract explicit olympiad/profile pairs from a structured table row."""
    if payload.get("kind") != "table_row":
        return []
    headers = payload.get("headers")
    row = payload.get("row")
    if not isinstance(headers, list) or not isinstance(row, str):
        return []
    header_line = next((item for item in headers if isinstance(item, str) and "|" in item), None)
    if header_line is None:
        return []
    header_cells = [normalized(cell) for cell in markdown_cells(header_line)]
    row_cells = markdown_cells(row)
    profile_index = next(
        (index for index, cell in enumerate(header_cells) if "профиль олимпиады" in cell),
        None,
    )
    if profile_index is None or len(row_cells) <= profile_index:
        return []
    profile = row_cells[profile_index]
    if not profile:
        return []
    title_index = olympiad_title_column_index(header_cells)
    title = row_cells[title_index] if title_index is not None and len(row_cells) > title_index else ""
    title = title or payload.get("inherited_olympiad_title")
    if isinstance(title, str) and title:
        return [{"olympiad_title": title, "profile_title": profile}]
    list_index = next(
        (index for index, cell in enumerate(header_cells) if "перечень олимпиад" in cell),
        None,
    )
    if list_index is None or len(row_cells) <= list_index:
        return []
    return [
        {"olympiad_title": title, "profile_title": profile}
        for title in explicit_olympiad_names(row_cells[list_index])
    ]


def table_header_index(table: list[str]) -> int:
    """Find a semantic header even if a PDF places a faculty row before it."""
    scored = [
        (sum(term in normalized(line) for term in TABLE_HEADER_TERMS), index)
        for index, line in enumerate(table)
        if "---" not in line
    ]
    score, index = max(scored, default=(0, 0))
    if score >= 2:
        return index
    # Ordinary HTML tables have a Markdown separator after their first row.
    separator = next((offset for offset, line in enumerate(table) if "---" in line), 0)
    return max(0, separator - 1)


def table_blocks(lines: list[str]) -> list[tuple[int, str, dict[str, Any]]]:
    """Return relevant Markdown table rows; headers become context for every row."""
    output: list[tuple[int, str, dict[str, Any]]] = []
    index = 0
    previous_heading = ""
    while index < len(lines):
        if lines[index].lstrip().startswith("#"):
            previous_heading = compact(lines[index].lstrip("# "))
        if not lines[index].lstrip().startswith("|"):
            index += 1
            continue
        start = index
        while index < len(lines) and lines[index].lstrip().startswith("|"):
            index += 1
        table = lines[start:index]
        header_index = table_header_index(table)
        headers = [table[header_index]]
        table_context = compact(" ".join([previous_heading, *table[header_index : header_index + 4]]))
        if not (inferred_benefit(table_context) or any(term in normalized(table_context) for term in CONTEXT_TERMS)):
            continue
        inherited_olympiad_title: str | None = None
        for offset, row in enumerate(table[header_index + 1 :], start=header_index + 1):
            if "---" in row:
                continue
            text = compact(" ".join([previous_heading, *headers, row]))
            if len(text) > 20:
                direct_title = explicit_row_olympiad_title(headers, row)
                if direct_title:
                    inherited_olympiad_title = direct_title
                payload: dict[str, Any] = {
                    "kind": "table_row",
                    "heading": previous_heading,
                    "headers": headers,
                    "row": row,
                }
                if inherited_olympiad_title and not direct_title:
                    payload["inherited_olympiad_title"] = inherited_olympiad_title
                output.append((start + offset + 1, text, payload))
    return output


def paragraph_blocks(lines: list[str]) -> list[tuple[int, str, dict[str, Any]]]:
    output: list[tuple[int, str, dict[str, Any]]] = []
    for index, line in enumerate(lines):
        if line.lstrip().startswith("|") or len(line) < 20:
            continue
        context = compact(" ".join(lines[max(0, index - 2) : min(len(lines), index + 3)]))
        if inferred_benefit(context) and any(term in normalized(context) for term in CONTEXT_TERMS):
            output.append((index + 1, context, {"kind": "paragraph", "line": line}))
    return output


def find_candidates(markdown: str, document_hash: str) -> list[KeywordCandidate]:
    lines = [compact(line) for line in markdown.splitlines() if compact(line)]
    seen: set[str] = set()
    result: list[KeywordCandidate] = []
    for line_number, excerpt, payload in [*table_blocks(lines), *paragraph_blocks(lines)]:
        excerpt = excerpt[:4_000]
        identities = table_identities(payload)
        for identity in identities or [None]:
            candidate_payload = {**payload, **({"identity": identity} if identity else {})}
            key = stable_key(
                document_hash,
                str(line_number),
                excerpt,
                identity["olympiad_title"] if identity else "",
                identity["profile_title"] if identity else "",
            )
            if key in seen:
                continue
            seen.add(key)
            benefit = inferred_benefit(excerpt)
            result.append(
                KeywordCandidate(
                    candidate_key=key,
                    source_locator=f"markdown:line:{line_number}",
                    source_excerpt=excerpt,
                    raw_rule_text=excerpt,
                    suggested_benefit_kind=benefit,
                    suggested_diploma_status=inferred_diploma(excerpt),
                    suggested_confirmation_subject_code=inferred_subject(excerpt),
                    suggested_confirmation_min_score=inferred_min_score(excerpt),
                    confidence=72 if payload["kind"] == "table_row" and benefit else 52,
                    raw_payload=candidate_payload,
                    raw_olympiad_name=identity["olympiad_title"] if identity else None,
                    raw_profile_name=identity["profile_title"] if identity else None,
                )
            )
    return result


def strict_rule_candidate(rule: StrictRule, document_hash: str) -> KeywordCandidate:
    """Make a strict table row auditable through the normal candidate table.

    A strict adapter has *not* published a benefit at this point.  The
    publisher still resolves the exact programme and RСОШ profile foreign keys
    inside the same admission campaign.
    """
    key = stable_key(
        document_hash,
        rule.source_locator,
        rule.programme_selector,
        rule.olympiad_title,
        rule.profile_title,
        rule.diploma_status,
        rule.benefit_kind,
    )
    return KeywordCandidate(
        candidate_key=key,
        source_locator=rule.source_locator,
        source_excerpt=rule.source_excerpt,
        raw_rule_text=rule.source_excerpt,
        suggested_benefit_kind=rule.benefit_kind,
        suggested_diploma_status=rule.diploma_status,
        suggested_confirmation_subject_code=rule.confirmation_subject_code,
        suggested_confirmation_min_score=rule.confirmation_min_score,
        confidence=100,
        raw_payload=rule.raw_payload,
        raw_olympiad_name=rule.olympiad_title,
        raw_profile_name=rule.profile_title,
        raw_programme_name=rule.programme_selector,
    )


def due_targets(
    connection: "psycopg.Connection",
    limit: int,
    *,
    force: bool = False,
    excluded_ids: tuple[int, ...] = (),
) -> list[Target]:
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT target.id, target.source_id, target.admission_campaign_id AS campaign_id,
                   target.university_location_id, target.adapter_code, target.url,
                   source.base_url AS source_base_url, target.adapter_config
            FROM admission_source_targets target
            JOIN sources source ON source.id = target.source_id
            WHERE target.is_active
              AND target.adapter_code = ANY(%s)
              AND (%s OR target.next_check_at IS NULL OR target.next_check_at <= now())
              AND NOT target.id = ANY(%s)
              AND NOT EXISTS (
                SELECT 1
                FROM admission_parse_runs running
                WHERE running.admission_source_target_id = target.id
                  AND running.status = 'running'
                  AND running.started_at > now() - interval '1 hour'
              )
            ORDER BY target.next_check_at NULLS FIRST, target.id
            LIMIT %s
            FOR UPDATE OF target SKIP LOCKED
            """,
            (list(SUPPORTED_TARGET_CODES), force, list(excluded_ids), limit),
        )
        return [Target(**row) for row in cursor.fetchall()]


def create_document_and_run(
    connection: "psycopg.Connection",
    target: Target,
    *,
    final_url: str,
    body: bytes,
    content_type: str,
    etag: str | None,
    last_modified: str | None,
    storage_key: str,
) -> tuple[int, int] | None:
    digest = hashlib.sha256(body).hexdigest()
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT document.id,
                   EXISTS (
                     SELECT 1 FROM admission_parse_runs run
                     WHERE run.admission_source_target_id = %s
                       AND run.source_document_id = document.id
                       AND run.adapter_version = %s
                       AND run.status = 'succeeded'
                   ) AS parsed
            FROM source_documents document
            WHERE document.source_id = %s AND document.content_sha256 = %s
            ORDER BY document.fetched_at DESC
            LIMIT 1
            """,
            (target.id, adapter_version(target), target.source_id, digest),
        )
        prior = cursor.fetchone()
        if prior and prior["parsed"]:
            cursor.execute(
                """
                UPDATE admission_source_targets
                SET last_checked_at = now(), next_check_at = now() + make_interval(hours => poll_interval_hours), updated_at = now()
                WHERE id = %s
                """,
                (target.id,),
            )
            connection.commit()
            return None
        if prior:
            document_id = prior["id"]
            # A failed conversion of the same immutable document is retryable.
            # Reuse its run because `(target, document, adapter_version)` is unique.
            cursor.execute(
                """
                SELECT id
                FROM admission_parse_runs
                WHERE admission_source_target_id = %s
                  AND source_document_id = %s
                  AND adapter_version = %s
                LIMIT 1
                """,
                (target.id, document_id, adapter_version(target)),
            )
            existing_run = cursor.fetchone()
            if existing_run:
                run_id = existing_run["id"]
                cursor.execute(
                    """
                    UPDATE admission_parse_runs
                    SET status = 'running', started_at = now(), finished_at = NULL,
                        records_seen = 0, candidates_created = 0, error_message = NULL
                    WHERE id = %s
                    """,
                    (run_id,),
                )
            else:
                run_id = None
        else:
            cursor.execute(
                """
                INSERT INTO source_documents (
                  source_id, url, content_sha256, storage_key, content_type, content_size_bytes,
                  http_etag, source_last_modified_at, title, metadata
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    target.source_id, final_url, digest, storage_key, content_type, len(body), etag,
                    parsed_http_date(last_modified), f"Admission document: {target.url}",
                    Jsonb({"crawler": target.adapter_code, "crawler_version": adapter_version(target)}),
                ),
            )
            document_id = cursor.fetchone()["id"]
            run_id = None
        if run_id is None:
            cursor.execute(
                """
                INSERT INTO admission_parse_runs (
                  admission_source_target_id, source_document_id, adapter_code, adapter_version,
                  status
                ) VALUES (%s, %s, %s, %s, 'running')
                RETURNING id
                """,
                (target.id, document_id, target.adapter_code, adapter_version(target)),
            )
            run_id = cursor.fetchone()["id"]
    connection.commit()
    return document_id, run_id


def finish_run(
    connection: "psycopg.Connection",
    target: Target,
    *,
    document_id: int,
    run_id: int,
    candidates: list[KeywordCandidate],
) -> int:
    with connection.cursor() as cursor:
        cursor.execute("SELECT id, code FROM subjects")
        subject_ids = {row["code"]: row["id"] for row in cursor.fetchall()}
        # A new parser version produces a fresh, auditable run for the same
        # immutable document. Its pending candidates replace older pending
        # candidates from that target; reviewed records are never changed.
        cursor.execute(
            """
            UPDATE admission_rule_candidates candidate
            SET review_status = 'superseded'
            WHERE candidate.review_status = 'pending'
              AND candidate.admission_parse_run_id IN (
                SELECT prior.id
                FROM admission_parse_runs prior
                WHERE prior.admission_source_target_id = %s
                  AND prior.id <> %s
              )
            """,
            (target.id, run_id),
        )
        for candidate in candidates:
            automatic_status = "pending_resolution" if target.adapter_code in strict_adapter_codes() else "not_applicable"
            cursor.execute(
                """
                INSERT INTO admission_rule_candidates (
                  admission_parse_run_id, source_document_id, admission_campaign_id, candidate_key,
                  source_locator, source_excerpt, raw_payload, raw_rule_text,
                  raw_programme_name, raw_olympiad_name, raw_profile_name,
                  suggested_diploma_status, suggested_benefit_kind,
                  suggested_point_value, suggested_olympiad_level,
                  suggested_confirmation_subject_id, suggested_confirmation_min_score,
                  match_status, review_status, automatic_status, automatic_note, confidence
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'unresolved', 'pending', %s, %s, %s)
                ON CONFLICT (admission_parse_run_id, candidate_key) DO NOTHING
                """,
                (
                    run_id, document_id, target.campaign_id, candidate.candidate_key,
                    candidate.source_locator, candidate.source_excerpt, Jsonb(candidate.raw_payload),
                    candidate.raw_rule_text, candidate.raw_programme_name,
                    candidate.raw_olympiad_name, candidate.raw_profile_name,
                    candidate.suggested_diploma_status,
                    candidate.suggested_benefit_kind,
                    candidate.suggested_point_value, candidate.suggested_olympiad_level,
                    subject_ids.get(candidate.suggested_confirmation_subject_code),
                    candidate.suggested_confirmation_min_score, automatic_status, None, candidate.confidence,
                ),
            )
        cursor.execute(
            """
            UPDATE admission_parse_runs
            SET status = 'succeeded', records_seen = %s, candidates_created = %s, finished_at = now()
            WHERE id = %s
            """,
            (len(candidates), len(candidates), run_id),
        )
        cursor.execute(
            """
            UPDATE admission_source_targets
            SET last_checked_at = now(), last_successful_document_id = %s,
                next_check_at = now() + make_interval(hours => poll_interval_hours), updated_at = now()
            WHERE id = %s
            """,
            (document_id, target.id),
        )
    connection.commit()
    return len(candidates)


def defer_target(connection: "psycopg.Connection", target_id: int) -> None:
    """Mark an attempted source as checked even when it could not be fetched.

    A transient DNS or HTTP failure must not make the same row due again in the
    current batch.  It will be retried on its normal polling schedule, while
    the worker can continue with the remaining universities.
    """
    with connection.cursor() as cursor:
        cursor.execute(
            """
            UPDATE admission_source_targets
            SET last_checked_at = now(),
                next_check_at = now() + make_interval(hours => poll_interval_hours),
                updated_at = now()
            WHERE id = %s
            """,
            (target_id,),
        )
    connection.commit()


def fail_run(connection: "psycopg.Connection", target: Target, run_id: int, message: str) -> None:
    with connection.cursor() as cursor:
        cursor.execute(
            "UPDATE admission_parse_runs SET status = 'failed', error_message = %s, finished_at = now() WHERE id = %s",
            (message[:4_000], run_id),
        )
        cursor.execute(
            """
            UPDATE admission_source_targets
            SET last_checked_at = now(),
                next_check_at = now() + make_interval(hours => poll_interval_hours),
                updated_at = now()
            WHERE id = %s
            """,
            (target.id,),
        )
    connection.commit()


def crawl_target(connection: "psycopg.Connection", target: Target) -> tuple[str, int]:
    body, final_url, content_type, etag, last_modified = fetch(target)
    digest = hashlib.sha256(body).hexdigest()
    suffix = guess_suffix(content_type, final_url)
    if suffix == ".bin":
        raise ValueError(f"Unsupported source content type: {content_type}")
    SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)
    snapshot = SNAPSHOT_DIR / f"{digest}{suffix}"
    if not snapshot.exists():
        snapshot.write_bytes(body)
    created = create_document_and_run(
        connection, target, final_url=final_url, body=body, content_type=content_type,
        etag=etag, last_modified=last_modified, storage_key=str(snapshot.relative_to(ROOT)),
    )
    if created is None:
        return "unchanged", 0
    document_id, run_id = created
    try:
        markdown = document_to_markdown(snapshot)
        candidates = (
            [strict_rule_candidate(rule, digest) for rule in strict_rules(target.adapter_config.get("strict_adapter", target.adapter_code), markdown)]
            if target.adapter_code in strict_adapter_codes()
            else find_candidates(markdown, digest)
        )
        return "succeeded", finish_run(connection, target, document_id=document_id, run_id=run_id, candidates=candidates)
    except Exception as error:
        fail_run(connection, target, run_id, str(error))
        raise


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--limit", type=int, default=10, help="maximum due official targets to crawl")
    parser.add_argument(
        "--force",
        action="store_true",
        help="reprocess active sources with the current parser version, ignoring their poll time",
    )
    return parser.parse_args()


def main() -> int:
    args = arguments()
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        print("DATABASE_URL is required", file=sys.stderr)
        return 2
    if psycopg is None or requests is None:
        print("Install dependencies first: python3 -m pip install -r requirements-crawler.txt", file=sys.stderr)
        return 2
    with psycopg.connect(database_url, row_factory=dict_row) as connection:
        errors = 0
        processed = 0
        processed_target_ids: set[int] = set()
        while processed < args.limit:
            # Claim one row only. The lock survives the HTTP request until the
            # parser run has been written; after that the `running` predicate
            # prevents another worker from claiming the same target.
            targets = due_targets(
                connection,
                1,
                force=args.force,
                excluded_ids=tuple(processed_target_ids) if args.force else (),
            )
            if not targets:
                if processed == 0:
                    print("No due official admission sources.")
                break
            target = targets[0]
            processed += 1
            processed_target_ids.add(target.id)
            try:
                status, count = crawl_target(connection, target)
                print(f"target {target.id}: {status}; {count} review candidates")
            except (requests.RequestException, OSError, RuntimeError, ValueError) as error:
                errors += 1
                defer_target(connection, target.id)
                print(f"target {target.id}: failed — {error}", file=sys.stderr)
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
